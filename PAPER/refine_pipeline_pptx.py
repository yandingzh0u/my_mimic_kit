"""Refine the existing editable pipeline; preserve a pre-edit backup."""
from pathlib import Path
import math
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement

ROOT = Path(__file__).resolve().parent / 'figures'
TARGET = ROOT / 'dare_pipeline_editable_fixed.pptx'
BACKUP = ROOT / 'dare_pipeline_editable_fixed_before_refine.pptx'
if not BACKUP.exists():
    shutil.copy2(TARGET, BACKUP)
p = Presentation(BACKUP)
s = p.slides[0]
old = list(s.shapes)
navy = '17365D'

def box(i, x, y, w, h):
    sh = old[i]
    sh.left, sh.top, sh.width, sh.height = map(Inches, (x,y,w,h))

for sh in old:
    if sh.has_text_frame:
        tf=sh.text_frame
        tf.margin_left=tf.margin_right=Inches(.015)
        tf.margin_top=tf.margin_bottom=0
        tf.vertical_anchor=MSO_ANCHOR.MIDDLE
        for q in tf.paragraphs:
            q.alignment=PP_ALIGN.CENTER
            q.space_before=q.space_after=Pt(0)
            for r in q.runs:
                r.font.name='Times New Roman'
box(27,5.6,1.43,3.26,.34)
box(65,9.55,1.43,3.4,.34)
box(63,9.36,1.38,3.78,3.82)
box(64,9.36,1.38,3.78,.45)
box(29,6.25,1.94,1.02,.44)
box(28,5.48,1.94,.76,.44)
box(30,7.24,1.94,.82,.44)
box(55,7.76,2.94,.66,.43)
box(56,7.76,3.53,.66,.26)
box(60,8.5,2.78,.6,.42)
box(8,.56,4.6,.95,.22)
box(221,10.82,5.87,1.7,.25)
box(225,5.92,6.08,1.7,.25)
# Correct the mathematical typography without changing the reward definition.
old[218].text='rₜ = β softplus(f(Δₜ))'
for r in old[218].text_frame.paragraphs[0].runs:
    r.font.name='Times New Roman';r.font.size=Pt(15)
old[218].text_frame.paragraphs[0].alignment=PP_ALIGN.CENTER

for i in range(67,211):
    el=old[i]._element;el.getparent().remove(el)

def text(t,x,y,w,h,size=12,bold=False):
    sh=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=sh.text_frame;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    q=tf.paragraphs[0];q.alignment=PP_ALIGN.CENTER
    r=q.add_run();r.text=t;r.font.name='Times New Roman';r.font.size=Pt(size)
    r.font.bold=bold;r.font.color.rgb=RGBColor.from_string(navy)
    return sh

def line(x1,y1,x2,y2,color='222222',width=1,arrow=False):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh.line.color.rgb=RGBColor.from_string(color);sh.line.width=Pt(width)
    if arrow:
        end=OxmlElement('a:tailEnd');end.set('type','triangle');end.set('w','sm');end.set('len','sm')
        sh.line._get_or_add_ln().append(end)
    return sh

def plot(x,y,w,h,points,directions=None):
    # All panels use the same softplus and score domain. The response band
    # represents sigmoid slopes 0.2--0.8: x in [-ln4, ln4].
    lo,hi=-2.7,2.7
    def xy(v):return (x+(v-lo)/(hi-lo)*w,y+h-math.log1p(math.exp(v))/math.log1p(math.exp(hi))*h*.91)
    left=xy(-math.log(4))[0];right=xy(math.log(4))[0]
    band=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(left),Inches(y),Inches(right-left),Inches(h))
    band.fill.solid();band.fill.fore_color.rgb=RGBColor.from_string('DCEED5');band.line.fill.background()
    line(x,y+h,x+w+.05,y+h,arrow=True)
    line(x,y+h,x,y-.025,arrow=True)
    coords=[xy(lo+(hi-lo)*i/120) for i in range(121)]
    path=s.shapes.build_freeform(Inches(coords[0][0]),Inches(coords[0][1]))
    path.add_line_segments([(Inches(a),Inches(b)) for a,b in coords[1:]],close=False)
    curve=path.convert_to_shape();curve.fill.background();curve.line.color.rgb=RGBColor.from_string('1678CA');curve.line.width=Pt(1.7)
    for v,c in zip(points,['E34A33','1678CA']):
        a,b=xy(v);d=.095
        dot=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(a-d/2),Inches(b-d/2),Inches(d),Inches(d))
        dot.fill.solid();dot.fill.fore_color.rgb=RGBColor.from_string(c);dot.line.color.rgb=RGBColor(255,255,255);dot.line.width=Pt(.5)
    if directions:
        for v,delta in zip(points,directions):
            a,b=xy(v);line(a+(.06 if delta>0 else -.06),b,a+delta,b,navy,1,True)
    text('score',x+w-.3,y+h+.045,.4,.17,10)

text('gap too large',9.53,1.94,1.52,.27,14,True)
text('gap too small',11.28,1.94,1.52,.27,14,True)
plot(9.62,2.35,1.27,.77,[-2.05,2.05],[.23,-.23])
plot(11.37,2.35,1.27,.77,[-.35,.35],[-.23,.23])
line(10.29,3.43,10.87,3.79,arrow=True)
line(12.00,3.43,11.43,3.79,arrow=True)
text('after calibration',10.43,3.81,1.47,.25,14,True)
plot(10.18,4.17,1.26,.73,[-math.log(4),math.log(4)])
box(218,9.92,5.31,2.62,.27)
# Center the reward-to-PPO connector over the PPO box.
old[219]._element.getparent().remove(old[219]._element)
line(11.67,5.6,11.67,5.78,'1678CA',1.5,True)
p.save(TARGET)
print(TARGET)
