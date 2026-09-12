"""Fix notation and draw exact softplus curves in dare_pipeline.pptx."""
from pathlib import Path
import math
import shutil

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent / "figures"
src = ROOT / "dare_pipeline.pptx"
bak = ROOT / "dare_pipeline_before_softplus.pptx"
if not bak.exists():
    shutil.copy2(src, bak)

p = Presentation(bak)
s = p.slides[0]

# Make the residual construction explicitly time-indexed.
s.shapes[17].text = "Δₜ = φ(ŝₜ) − φ(sₜ)"
for run in s.shapes[17].text_frame.paragraphs[0].runs:
    run.font.name = "Times New Roman"

# The old curves (84, 93, 105) were quadratic freeform approximations.
old_curves = [s.shapes[i] for i in (84, 93, 105)]
dot_anchors = [s.shapes[85], s.shapes[94], s.shapes[106]]
for sh in old_curves:
    el = sh._element
    el.getparent().remove(el)

def softplus_curve(x, y, w, h):
    # Score domain and vertical normalization are shared by all three panels.
    lo, hi = -2.7, 2.7
    base = y + h
    scale = math.log1p(math.exp(hi))
    pts = []
    for i in range(161):
        z = lo + (hi - lo) * i / 160.0
        val = math.log1p(math.exp(z))
        pts.append((x + (z - lo) / (hi - lo) * w,
                    base - val / scale * h))
    path = s.shapes.build_freeform(Inches(pts[0][0]), Inches(pts[0][1]))
    path.add_line_segments([(Inches(a), Inches(b)) for a, b in pts[1:]], close=False)
    curve = path.convert_to_shape()
    curve.fill.background()
    curve.line.color.rgb = RGBColor.from_string("1678CA")
    curve.line.width = Pt(1.7)
    return curve

boxes = [(9.62, 2.42, 1.27, .68),
         (11.37, 2.42, 1.27, .68),
         (10.57, 4.24, 1.26, .65)]
for box, dot in zip(boxes, dot_anchors):
    curve = softplus_curve(*box)
    # Keep the curve behind the marker dots, as in a normal plotted figure.
    dot._element.addprevious(curve._element)

p.save(src)
print(src)
